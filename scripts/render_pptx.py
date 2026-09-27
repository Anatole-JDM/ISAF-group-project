"""Render the deck to .pptx with python-pptx, from scripts/deck_content.py.

Same content and same geometry as scripts/render_pdf.py, so the two deliverables
carry identical text and identical numbers. Speaker notes go in the notes pane,
which is where a six-person group needs them.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import deck_content as C

FIG = ROOT / "outputs" / "deck"

INK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)
BLUE = RGBColor(0x2C, 0x7F, 0xB8)
GREY = RGBColor(0x8A, 0x89, 0x84)
PALE = RGBColor(0xF2, 0xF1, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0xC9, 0xC8, 0xD4)
FONT = "Arial"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.72)
CW = W - 2 * M


def _tb(slide, l, t, w, h):
    tf = slide.shapes.add_textbox(l, t, w, h).text_frame
    tf.word_wrap = True
    return tf


def _run(p, text, size, *, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    return r


def _para(tf, text="", size=16, *, bold=False, color=INK, space_after=8, italic=False,
          first=False, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if text:
        _run(p, text, size, bold=bold, color=color, italic=italic)
    return p


def _rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def draw_stats(s, stats, top):
    n = len(stats)
    height = Inches(1.42) if n == 3 else Inches(1.5)
    gap = Inches(0.22)
    bw = int((CW - gap * (n - 1)) / n)
    for i, (val, label) in enumerate(stats):
        l = M + i * (bw + gap)
        _rect(s, l, top, bw, height, PALE)
        tf = _tb(s, l + Inches(0.14), top + Inches(0.12), bw - Inches(0.28), height - Inches(0.24))
        _para(tf, val, 30, bold=True, color=ACCENT if (i == 2 and n == 3) else INK,
              first=True, space_after=2, align=PP_ALIGN.CENTER)
        for j, ln in enumerate(label.split("\n")):
            _para(tf, ln, 10.5, color=INK, space_after=0, align=PP_ALIGN.CENTER)
    return top + height


def draw_bullets(s, items, top, size=15.0):
    tf = _tb(s, M, top, CW, H - top - Inches(0.6))
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            head, body = it
            p = _para(tf, size=size, space_after=3, first=(i == 0))
            _run(p, head, size, bold=True)
            _para(tf, body, size - 2.5, color=GREY, space_after=13)
        else:
            p = _para(tf, size=size, space_after=11, first=(i == 0))
            _run(p, "— ", size, color=ACCENT, bold=True)
            _run(p, it, size)
    return top


def draw_picture(s, name, top, height_in):
    pic = s.shapes.add_picture(str(FIG / name), 0, top, height=Inches(height_in))
    if pic.width > CW:
        ratio = CW / pic.width
        pic.width, pic.height = int(CW), int(pic.height * ratio)
    pic.left = int((W - pic.width) / 2)
    return top + pic.height


def draw_table(s, spec, top, size=11.5):
    headers, rows = spec["headers"], spec["rows"]
    cw = spec.get("col_w") or [1] * len(headers)
    total = sum(cw)
    tbl = s.shapes.add_table(len(rows) + 1, len(headers), M, top, CW,
                             Inches(0.34) * (len(rows) + 1)).table
    for i, v in enumerate(cw):
        tbl.columns[i].width = Emu(int(CW * v / total))
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = ""
        c.fill.solid(); c.fill.fore_color.rgb = INK
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
        _run(p, h, size, bold=True, color=WHITE)
    hl_neg, hl_row = spec.get("highlight_neg"), spec.get("highlight_row")
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            v = str(v)
            c = tbl.cell(i, j)
            c.text = ""
            c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else PALE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
            neg = hl_neg is not None and j == hl_neg and v.lstrip().startswith(("−", "-"))
            hot = neg or (hl_row is not None and i - 1 == hl_row and j >= 2) or v == "SIGNIFICANT"
            _run(p, v, size, bold=(hot or j == 0), color=ACCENT if hot else INK)
    return top + Inches(0.34) * (len(rows) + 1)


def draw_callout(s, text, *, color=ACCENT, height=None):
    height = height or Inches(0.95)
    top = H - Inches(0.45) - height
    _rect(s, M, top, CW, height, PALE)
    _rect(s, M, top, Pt(4.5), height, color)
    tf = _tb(s, M + Inches(0.22), top + Inches(0.1), CW - Inches(0.44), height - Inches(0.2))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, text, 14, bold=True, color=INK, first=True, space_after=0)


def draw_caption(s, text, top, *, color=GREY):
    tf = _tb(s, M, top, CW, Inches(0.7))
    _para(tf, text, 11, color=color, italic=True, first=True, space_after=0)


def build(dest: Path | None = None) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    for idx, spec in enumerate(C.SLIDES, start=1):
        s = prs.slides.add_slide(prs.slide_layouts[6])

        if spec.get("cover"):
            _rect(s, 0, 0, W, H, INK)
            _rect(s, M, Inches(2.42), Inches(2.0), Pt(5), ACCENT)
            tf = _tb(s, M, Inches(2.72), Inches(11.0), Inches(1.9))
            for i, ln in enumerate(spec["headline"]):
                _para(tf, ln, 38, bold=True, color=WHITE, first=(i == 0), space_after=4)
            _para(tf, spec["sub"], 16, color=DIM, space_after=0)
            tf = _tb(s, M, H - Inches(1.5), Inches(11.5), Inches(1.0))
            for i, ln in enumerate(spec["foot"]):
                _para(tf, ln, 12, color=GREY, first=(i == 0), space_after=3)
            s.notes_slide.notes_text_frame.text = spec.get("notes", "")
            continue

        top = M
        if spec.get("kicker"):
            tf = _tb(s, M, Inches(0.34), CW, Inches(0.32))
            _para(tf, spec["kicker"].upper(), 11, bold=True, color=ACCENT, first=True, space_after=0)
            top = Inches(0.72)
        if spec.get("title"):
            tf = _tb(s, M, top, CW, Inches(0.9))
            _para(tf, spec["title"], 29, bold=True, first=True, space_after=0)
            _rect(s, M, top + Inches(0.82), Inches(1.5), Pt(3.2), ACCENT)
            top += Inches(1.18)

        body = top
        if "stats" in spec:
            body = draw_stats(s, spec["stats"], body) + Inches(0.28)
        if "picture" in spec:
            body = draw_picture(s, spec["picture"], body, spec.get("picture_h", 4.0)) + Inches(0.12)
        if "table" in spec:
            body = draw_table(s, spec["table"], body) + Inches(0.16)
        if "bullets" in spec:
            body = draw_bullets(s, spec["bullets"], body, size=spec.get("bullet_size", 15.0))
        if "caption" in spec:
            draw_caption(s, spec["caption"], body if ("table" in spec) else H - Inches(1.12),
                         color=INK if "table" in spec else GREY)
        if "callout" in spec:
            draw_callout(s, spec["callout"],
                         color=BLUE if spec.get("callout_color") == "blue" else ACCENT,
                         height=Inches(spec["callout_h"]) if spec.get("callout_h") else None)

        tf = _tb(s, W - Inches(1.15), H - Inches(0.46), Inches(0.62), Inches(0.3))
        _para(tf, str(idx), 10.5, color=GREY, first=True, space_after=0, align=PP_ALIGN.RIGHT)
        s.notes_slide.notes_text_frame.text = spec.get("notes", "")

    dest = dest or (ROOT / "reports" / "ISAF_presentation.pptx")
    dest.parent.mkdir(exist_ok=True)
    prs.save(dest)
    print(f"wrote {dest.relative_to(ROOT)}  ({len(C.SLIDES)} slides)")
    return dest


if __name__ == "__main__":
    build()
